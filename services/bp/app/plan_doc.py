"""Compose a readable business plan from the founder's answers, and export it as a Word file.

Nothing here invents content. Each section is the founder's own words under a clear sub-heading,
plus figures computed from the numbers they gave (tables, totals, projection). Anything missing is listed
under "To complete" and in an open-items page, never filled with a guess.
With a model key, ``polish`` may rewrite each section as plan prose. Three guards then run on the result:
the figure guard (every number not in the inputs is marked), the claim guard (a sentence that asserts
experience, a source, a market size, a partner or a pilot the founder never mentioned becomes a visible
[NEEDED: ...] question, and a foreign currency is rewritten), and a consistency pass over the answers.
With the final writer's six sections available, the plan opens with that narrative.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Inches, Pt

from app.agents import claim_guard
from app.agents.writer import MARKERS, _fmt  # noqa: F401
from app.questionnaire import BANK_BY_ID
from app.templates import label

# confirm markers and [NEEDED: ...] placeholders are highlighted in the Word file
HILITE = re.compile(r"\[(?:to be confirmed by the founder|à confirmer par le fondateur|يؤكده الباعث)\]|\[NEEDED:[^\]]*\]", re.I)
_UNKNOWN_TOKEN = re.compile(r"\[\[[A-Z]+_\d+\]\]")
_SCORE_RE = re.compile(r"\s*\((?:score\s*)?\d{1,3}\s*/\s*100\)", re.I)

SUB = {
    "problem": {"en": "The problem", "fr": "Le problème", "ar": "المشكل"},
    "solution": {"en": "The offer", "fr": "L'offre", "ar": "العرض"},
    "differentiator": {"en": "What makes it different", "fr": "Ce qui nous différencie", "ar": "ما يميزنا"},
    "target_customers": {"en": "Target customers", "fr": "Clientèle cible", "ar": "الحرفاء المستهدفون"},
    "market_size_note": {"en": "Market size", "fr": "Taille du marché", "ar": "حجم السوق"},
    "competitors": {"en": "Competitors and alternatives", "fr": "Concurrents et alternatives", "ar": "المنافسون والبدائل"},
    "pricing": {"en": "Pricing", "fr": "Prix", "ar": "الأسعار"},
    "channels": {"en": "Sales channels", "fr": "Canaux de vente", "ar": "قنوات البيع"},
    "founders": {"en": "Founders and roles", "fr": "Promoteurs et rôles", "ar": "الباعثون وأدوارهم"},
    "key_skills": {"en": "Experience and skills", "fr": "Expérience et compétences", "ar": "الخبرة والمهارات"},
    "production_means": {"en": "Equipment and premises", "fr": "Équipements et locaux", "ar": "التجهيزات والمحلات"},
    "suppliers": {"en": "Suppliers", "fr": "Fournisseurs", "ar": "المزودون"},
    "risks": {"en": "Main risks and how we handle them", "fr": "Principaux risques et mesures", "ar": "المخاطر وسبل معالجتها"},
    "traction_note": {"en": "Traction and validation", "fr": "Traction et validation", "ar": "الإثبات والتحقق"},
    "ip_note": {"en": "Intellectual property", "fr": "Propriété intellectuelle", "ar": "الملكية الفكرية"},
    "awards_note": {"en": "Awards and support received", "fr": "Prix et soutiens obtenus", "ar": "الجوائز والدعم الحاصل عليه"},
    "permits_note": {"en": "Permits and regulatory requirements", "fr": "Autorisations et exigences réglementaires", "ar": "التراخيص والمتطلبات التنظيمية"},
    "extension_goal": {"en": "The extension", "fr": "L'extension", "ar": "التوسعة"},
}
TITLES = {"cover_sub": {"en": "Business plan: working draft", "fr": "Business plan : version de travail", "ar": "خطة العمل: مسودة"},
          "banner": {"en": "Draft generated from your answers. Yellow markers show what is still missing or must be confirmed by you. It is not a final or approved plan.",
                     "fr": "Version générée à partir de vos réponses. Les repères jaunes indiquent ce qui manque ou doit être confirmé par vous. Ce n'est pas un plan final ni validé.",
                     "ar": "مسودة أنشئت من إجاباتكم. العلامات الصفراء تبين ما ينقص أو يجب أن تؤكدوه. ليست خطة نهائية."},
          "key": {"en": "Key figures", "fr": "Chiffres clés", "ar": "أرقام أساسية"},
          "todo": {"en": "To complete", "fr": "À compléter", "ar": "للاستكمال"},
          "open": {"en": "Open items", "fr": "Points à compléter", "ar": "نقاط مفتوحة"},
          "uses": {"en": "Uses", "fr": "Emplois", "ar": "الاستخدامات"}, "sources": {"en": "Sources", "fr": "Ressources", "ar": "الموارد"},
          "proj": {"en": "Projection (computed from your figures)", "fr": "Projection (calculée à partir de vos chiffres)", "ar": "التوقعات"},
          "total": {"en": "Total", "fr": "Total", "ar": "المجموع"},
          "none_listed": {"en": "No detail was given.", "fr": "Aucun détail fourni.", "ar": "لا تفاصيل."},
          "narrative": {"en": "Plan narrative", "fr": "Synthèse du plan (en anglais)", "ar": "ملخص الخطة (بالإنجليزية)"},
          "verdict": {"en": "Self-assessment", "fr": "Auto-évaluation", "ar": "تقييم ذاتي"},
          "details": {"en": "Detailed answers and figures", "fr": "Réponses détaillées et chiffres", "ar": "الإجابات التفصيلية والأرقام"},
          "final_open": {"en": "Questions raised by the narrative", "fr": "Questions soulevées par la synthèse", "ar": "أسئلة أثارها الملخص"},
          "consistency": {"en": "Consistency checks", "fr": "Contrôles de cohérence", "ar": "فحوصات الاتساق"}}


def _t(key: str, lang: str) -> str:
    return TITLES[key].get(lang) or TITLES[key]["en"]


def _qlabel(qid: str, lang: str) -> str:
    if qid in SUB:
        return SUB[qid].get(lang) or SUB[qid]["en"]
    q = BANK_BY_ID[qid]
    return {"en": q.en, "fr": q.fr, "ar": q.ar}.get(lang, q.en)


def _num(v: float) -> str:
    return "-" if v is None else f"{v:,.0f}".replace(",", " ")


def _tidy(s: str) -> str:
    """Founder text as typed, minus the typing slips: repeated spaces, a doubled full stop, a lower-case first letter."""
    s = re.sub(r"[ \t]+", " ", str(s)).strip()
    s = re.sub(r"(?<!\.)\.\.(?!\.)", ".", s)
    m = re.match(r"[a-zà-ÿ]{3,}\b", s)
    if m:
        s = s[0].upper() + s[1:]
    return s


def _para_with_markers_text(doc, text: str) -> None:
    """Add paragraphs to a python-docx document, highlighting the confirm markers and the [NEEDED: ...] questions."""
    from docx.enum.text import WD_COLOR_INDEX

    for line in text.split("\n"):
        p = doc.add_paragraph()
        pos = 0
        for m in HILITE.finditer(line):
            p.add_run(line[pos:m.start()])
            r = p.add_run(m.group(0))
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            pos = m.end()
        p.add_run(line[pos:])


def compose_section(sec: dict, values: dict, derived: dict | None, lang: str) -> dict:
    """Return {id, title, blocks, missing}. A block is ('h', text) | ('p', text) | ('ul', [text]) | ('table', header, rows)."""
    marker = MARKERS.get(lang, MARKERS["en"])
    cur = str(values.get("currency") or "TND")
    blocks: list[tuple] = []
    missing: list[str] = []
    figures: list[str] = []
    for qid in sec["questions"]:
        q = BANK_BY_ID.get(qid)
        if q is None or q.type == "year_series":
            continue
        if qid not in values:
            if q.required or sec["id"] in ("risks",):
                missing.append(_qlabel(qid, lang))
            continue
        v = values[qid]
        if q.type in ("text", "longtext"):
            if qid == "competitors" and isinstance(v, str):
                parts = [p.strip() for p in re.split(r"[,;\n]+", v) if p.strip()]
                if len(parts) >= 2:
                    blocks += [("h", _qlabel(qid, lang)), ("ul", parts)]
                    continue
            blocks += [("h", _qlabel(qid, lang)), ("p", _tidy(v))]
        elif q.type == "items":
            continue  # shown as a table below
        elif q.type == "bool":
            figures.append(f"{_qlabel(qid, lang)}: {'yes' if v else 'no'}")
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            unit = f" {cur}" if q.type == "number" and ("TND" in q.en or cur != "TND") else (" %" if q.type == "pct" else "")
            figures.append(f"{_qlabel(qid, lang)}: {_num(v)}{unit}" if abs(v) >= 1000 else f"{_qlabel(qid, lang)}: {v:g}{unit}")
        else:
            figures.append(f"{_qlabel(qid, lang)}: {v}")
    if figures:
        blocks.append(("h", _t("key", lang)))
        blocks.append(("ul", figures))

    sid = sec["id"]
    if sid == "investment_plan" and values.get("capex_items"):
        rows = [[i["label"], _num(i["amount"])] for i in values["capex_items"]]
        rows.append([_t("total", lang), _num(sum(i["amount"] for i in values["capex_items"]))])
        blocks.append(("table", [_t("uses", lang), cur], rows))
    if sid == "financing_plan":
        src = [(k, values[k]) for k in ("own_funds", "loan_amount", "grant_amount") if values.get(k)]
        if src:
            rows = [[_qlabel(k, lang), _num(v)] for k, v in src]
            rows.append([_t("total", lang), _num(sum(v for _, v in src))])
            blocks.append(("table", [_t("sources", lang), cur], rows))
            ti = values.get("total_investment")
            gap = (ti - sum(v for _, v in src)) if ti else None
            if gap and abs(gap) > 1:
                blocks.append(("p", f"{marker} " + ({"en": f"Total investment is {_num(ti)} {cur} but the listed sources add up to {_num(sum(v for _, v in src))} {cur}.",
                                                     "fr": f"L'investissement total est de {_num(ti)} {cur} mais les ressources listées font {_num(sum(v for _, v in src))} {cur}.",
                                                     "ar": ""}[lang] if lang in ("en", "fr") else "")))
    if sid == "financial_projections" and derived and derived.get("revenue") and any(derived["revenue"]):
        yrs = derived["years"]
        incomplete = bool(derived.get("incomplete"))
        rows = []
        for key, lab in ((("revenue", {"en": "Revenue", "fr": "Chiffre d'affaires"}),) if incomplete else
                         (("revenue", {"en": "Revenue", "fr": "Chiffre d'affaires"}), ("ebitda", {"en": "EBITDA", "fr": "EBE"}),
                          ("net_income", {"en": "Net result", "fr": "Résultat net"}), ("cash_end", {"en": "Year-end cash", "fr": "Trésorerie fin d'année"}))):
            rows.append([lab.get(lang, lab["en"])] + [_num(x) for x in derived[key]])
        if incomplete and any(derived.get("cogs") or []):
            margin = [(r - c) if r is not None and c is not None else None for r, c in zip(derived["revenue"], derived["cogs"])]
            rows.append([{"en": "Direct margin", "fr": "Marge directe"}.get(lang, "Direct margin")] + [_num(x) for x in margin])
        blocks.append(("h", _t("proj", lang)))
        blocks.append(("table", [""] + list(yrs), rows))
        if incomplete:
            blocks.append(("p", {"en": "Projection incomplete: salaries or other costs were not provided, so only revenue and direct margin are shown.",
                                 "fr": "Projection incomplète : salaires ou autres charges non fournis, seuls le chiffre d'affaires et la marge directe sont présentés."}[lang if lang in ("en", "fr") else "en"]))
        else:
            be = derived.get("break_even_year")
            blocks.append(("p", ({"en": f"Break-even is reached in {be}.", "fr": f"Le point mort est atteint en {be}."} if be else
                                 {"en": "Break-even is not reached within the projection horizon.", "fr": "Le point mort n'est pas atteint sur l'horizon de projection."})[lang if lang in ("en", "fr") else "en"]))
        for a in derived.get("assumptions", []):
            blocks.append(("p", a))
        blocks.append(("chart",))
    if sid == "financial_projections" and not (derived and derived.get("revenue") and any(derived["revenue"])):
        missing.append({"en": "Revenue per year", "fr": "Chiffre d'affaires par année", "ar": "رقم المعاملات"}[lang if lang in ("en", "fr", "ar") else "en"])
    if missing:
        blocks.append(("h", _t("todo", lang)))
        blocks.append(("ul", [f"{m} {marker}" for m in missing]))
    return {"id": sid, "title": label(sec["title"], lang), "blocks": blocks, "missing": missing}


def section_plaintext(composed: dict) -> str:
    """Flatten composed blocks to plain text for critic quote matching."""
    parts: list[str] = []
    for b in composed.get("blocks") or []:
        kind = b[0]
        if kind == "p":
            parts.append(str(b[1]))
        elif kind == "ul":
            parts.extend(str(x) for x in b[1])
        elif kind == "table":
            for row in b[2]:
                parts.append(" ".join(str(c) for c in row))
    return "\n".join(parts)


def critic_context(section_id: str, sections: dict, composed: dict, full_text: str) -> str:
    """Text the critic may quote for this section: upload segment, composed body, and flat uploads."""
    chunks = [
        (sections.get(section_id) or {}).get("text", ""),
        section_plaintext(composed),
    ]
    if full_text and section_id in ("market_analysis", "financial_projections", "financing_plan", "project_description", "executive_summary"):
        chunks.append(full_text)
    return "\n".join(c for c in chunks if c).strip()


def build_executive_summary(composed: dict, sec: dict, values: dict, derived: dict | None, lang: str) -> None:
    """Synthesise the executive summary last from verified facts only (no new claims)."""
    marker = MARKERS.get(lang, MARKERS["en"])
    cur = str(values.get("currency") or "TND")
    fr = lang == "fr"
    lines: list[str] = []
    name, sector, stage = values.get("project_name"), values.get("sector"), values.get("stage")
    if name and sector:
        lines.append(f"{name} : {sector}." if fr else f"{name}: {sector}.")
    elif name:
        lines.append(f"{name}.")
    if stage:
        lines.append({"creation": "Projet en phase de création.", "extension": "Projet d'extension.", "idea": "Projet à l'étape idée."}.get(stage, stage) if fr
                       else f"Stage: {stage}.")
    if values.get("problem"):
        lines.append(_tidy(values["problem"]).rstrip(".") + ".")
    if values.get("solution"):
        lines.append(_tidy(values["solution"]).rstrip(".") + ".")
    if values.get("target_customers"):
        lines.append((f"Clientèle cible : {_tidy(values['target_customers'])}." if fr
                      else f"Target customers: {_tidy(values['target_customers']).rstrip('.')}."))
    if values.get("traction_note"):
        lines.append(_tidy(values["traction_note"]).rstrip(".") + ".")
    if values.get("pricing"):
        lines.append((f"Prix : {values['pricing']}." if fr else f"Pricing: {values['pricing']}."))
    else:
        lines.append((f"Prix : {marker}" if fr else f"Pricing: {marker}"))
    if values.get("funding_need_amount"):
        use = ""
        if values.get("capex_items"):
            use = (" ; emplois : " if fr else "; uses: ") + "; ".join(f"{i['label']} {_num(i['amount'])} {cur}" for i in values["capex_items"])
        lines.append((f"Financement externe recherché : {_num(values['funding_need_amount'])} {cur}{use}." if fr
                        else f"External funding sought: {_num(values['funding_need_amount'])} {cur}{use}."))
    elif values.get("total_investment"):
        lines.append((f"Investissement total : {_num(values['total_investment'])} {cur}." if fr
                      else f"Total investment: {_num(values['total_investment'])} {cur}."))
    if derived and derived.get("revenue") and any(derived["revenue"]):
        yrs, rev = derived["years"], derived["revenue"]
        traj = ", ".join(f"{y} {_num(v)} {cur}" for y, v in zip(yrs, rev) if v is not None)
        lines.append((f"Trajectoire de chiffre d'affaires : {traj}." if fr else f"Revenue trajectory: {traj}."))
        if derived.get("incomplete"):
            lines.append({"en": "Financial projection is incomplete (operating costs missing); profit and cash are not stated here.",
                          "fr": "Projection financière incomplète (charges d'exploitation manquantes) ; résultat et trésorerie non indiqués ici."}[lang if lang in ("en", "fr") else "en"])
    missing = composed.get("missing") or []
    if missing:
        lines.append((f"Points à compléter : {', '.join(missing[:6])}." if fr else f"Items still to complete: {', '.join(missing[:6])}."))
    title = label(sec["title"], lang)
    composed["blocks"] = [("p", " ".join(lines))] if lines else composed["blocks"]
    composed["title"] = title


def compose_summary(values: dict, derived: dict | None, lang: str, polished: bool = False) -> list[tuple]:
    """A short summary built only from answers that exist.

    When the executive summary has been rewritten as prose (``polished``) the rewrite already says what the business
    does and for whom, so only the money lines are added: no sentence is written twice.
    """
    name = values.get("project_name")
    cur = str(values.get("currency") or "TND")
    parts: list[str] = []
    fr = lang == "fr"
    if not polished:
        sector = _tidy(values["sector"]).rstrip(".") if values.get("sector") else None
        if name and sector:
            parts.append((f"{name} : {sector}." if fr else f"{name}: {sector}."))
        elif name:
            parts.append(name + ".")
        if values.get("solution"):
            parts.append(_tidy(values["solution"]).rstrip(".") + ".")
        if values.get("target_customers"):
            parts.append((f"Clientèle visée : {_tidy(values['target_customers'])}." if fr else f"Target customers: {_tidy(values['target_customers']).rstrip('.')}."))
    if values.get("total_investment"):
        extra = ""
        if values.get("funding_need_amount"):
            extra = (f" dont {_num(values['funding_need_amount'])} {cur} recherchés à l'extérieur" if fr else f", of which {_num(values['funding_need_amount'])} {cur} is sought externally")
        parts.append((f"Investissement total : {_num(values['total_investment'])} {cur}{extra}." if fr else f"Total investment: {_num(values['total_investment'])} {cur}{extra}."))
    elif values.get("funding_need_amount"):
        parts.append((f"Financement externe recherché : {_num(values['funding_need_amount'])} {cur}." if fr
                        else f"External funding sought: {_num(values['funding_need_amount'])} {cur}."))
    if derived and derived.get("revenue") and any(derived["revenue"]):
        last = max(i for i, v in enumerate(derived["revenue"]) if v)
        parts.append((f"Chiffre d'affaires prévu en {derived['years'][last]} : {_num(derived['revenue'][last])} {cur}." if fr
                      else f"Revenue expected in {derived['years'][last]}: {_num(derived['revenue'][last])} {cur}."))
    return [("p", " ".join(parts))] if parts else []


# ---------------------------------------------------------------------------- consistency of the answers
_ROLE_WORDS = re.compile(r"\b(?:ceo|cto|cfo|coo|cmo|co-?founder|founder|fondateur|fondatrice|gérant|gerant|directeur|directrice|pdg|dg)\b", re.I)


def consistency_notes(values: dict, lang: str) -> list[str]:
    """Contradictions inside the founder's own answers. Each note is a question to settle, never a correction."""
    fr = lang == "fr"
    notes: list[str] = []

    def say(en: str, fr_text: str) -> None:
        notes.append(fr_text if fr else en)

    need, total = values.get("funding_need_amount"), values.get("total_investment")
    cur = str(values.get("currency") or "TND")
    if need and total and need > total * 1.001:
        say(f"External funding sought ({_num(need)} {cur}) is larger than the total investment ({_num(total)} {cur}).",
            f"Le financement externe recherché ({_num(need)} {cur}) dépasse l'investissement total ({_num(total)} {cur}).")
    instrument = values.get("funding_instrument")
    if instrument == "grant" and values.get("loan_amount"):
        say("The main funding instrument is a grant, yet a bank loan is also listed. Which one is the main source?",
            "L'instrument principal est une subvention, mais un crédit bancaire est aussi indiqué. Laquelle est la source principale ?")
    if instrument == "none" and need:
        say("No outside funding is planned, yet an external funding need is stated.",
            "Aucun financement externe n'est prévu, mais un besoin de financement externe est indiqué.")
    if instrument in ("loan", "mixed") and not values.get("loan_amount") and not need:
        say("A loan is the funding route, but neither a loan amount nor an external funding need is given.",
            "Le crédit est la voie de financement, mais ni son montant ni le besoin externe ne sont indiqués.")
    roles = {m.group(0).lower().replace("-", "") for m in _ROLE_WORDS.finditer(str(values.get("founders") or ""))}
    roles = {"cofounder" if r.startswith("co") and "founder" in r else r for r in roles}
    head = values.get("headcount")
    if head and len(roles) > head:
        say(f"The team description names {len(roles)} roles but the headcount at start is {head}. Who is paid from day one?",
            f"La description de l'équipe cite {len(roles)} rôles mais l'effectif au démarrage est de {head}. Qui est rémunéré dès le premier jour ?")
    return notes


# ---------------------------------------------------------------------------- Word file
def _usable_narrative(final_plan: dict | None) -> list[dict]:
    return [s for s in (final_plan or {}).get("sections", []) if s.get("status") == "ok" and str(s.get("text", "")).strip()]


def _clean_why(why: str) -> str:
    """The critic's internal wording, minus the score: the founder needs the point, not the metric."""
    return _SCORE_RE.sub("", why or "").replace("is present but thin", "needs more detail").strip()


def build_plan_docx(title: str, sections: list[dict], chart_png: bytes | None, lang: str, open_items: list[tuple[str, list[str]]],
                    critique: dict | None = None, final_plan: dict | None = None) -> bytes:
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(title or {"en": "Business plan", "fr": "Business plan", "ar": "خطة العمل"}[lang], 0)
    d.add_paragraph(_t("cover_sub", lang))
    p = d.add_paragraph()
    p.add_run(_t("banner", lang)).italic = True
    # one disclaimer at the top of the document, not one note per section
    n_polished = sum(1 for s in sections if s.get("polished"))
    n_removed = sum(1 for s in sections for r in s.get("claims_removed", []) if r.get("rule") not in ("currency", "hype"))
    if n_polished or n_removed:
        note = POLISH_NOTE.get(lang, POLISH_NOTE["en"])
        if n_removed:
            note += REMOVED_NOTE.get(lang, REMOVED_NOTE["en"]).format(n=n_removed)
        p = d.add_paragraph()
        p.add_run(note).italic = True
    narrative = _usable_narrative(final_plan)
    if narrative:
        d.add_heading(_t("narrative", lang), 1)
        verdict = str((final_plan or {}).get("verdict") or "").strip()
        if verdict:
            p = d.add_paragraph()
            p.add_run(f"{_t('verdict', lang)}: {verdict}").italic = True
        for s in narrative:
            d.add_heading(s["title"], 2)
            _para_with_markers_text(d, s["text"])
        d.add_heading(_t("details", lang), 1)
    for sec in sections:
        d.add_heading(sec["title"], 1)
        for b in sec["blocks"]:
            kind = b[0]
            if kind == "h":
                d.add_heading(b[1], 2)
            elif kind == "p":
                _para_with_markers_text(d, b[1])
            elif kind == "ul":
                for item in b[1]:
                    _bullet(d, item)
            elif kind == "table":
                t = d.add_table(rows=1, cols=len(b[1]))
                t.style = "Light Grid Accent 1"
                for i, h in enumerate(b[1]):
                    t.rows[0].cells[i].text = h
                for r in b[2]:
                    cells = t.add_row().cells
                    for i, v in enumerate(r):
                        cells[i].text = str(v)
                d.add_paragraph()
            elif kind == "chart" and chart_png:
                d.add_picture(io.BytesIO(chart_png), width=Inches(6.0))
    final_open = [str(x) for x in (final_plan or {}).get("open_items", []) if str(x).strip()] if narrative else []
    if open_items or final_open:
        d.add_page_break()
        d.add_heading(_t("open", lang), 1)
        for sec_title, items in open_items:
            d.add_heading(sec_title, 2)
            for it in items:
                _bullet(d, it)
        if final_open:
            d.add_heading(_t("final_open", lang), 2)
            for it in final_open:
                _bullet(d, it)
    cw = (critique or {}).get("weaknesses") or []
    if cw:
        d.add_page_break()
        d.add_heading({"fr": "Faiblesses critiques à corriger", "ar": "نقاط ضعف حرجة يجب تصحيحها"}.get(lang, "Critical weaknesses to fix"), 1)
        fix_h = {"fr": "Correction", "ar": "الإصلاح"}.get(lang, "Fix")
        from app.agents.critic import dim_label
        for w in cw:
            _bullet(d, f"{dim_label(w['dimension'], lang)} — {_clean_why(w['why'])}")
            p = d.add_paragraph(f"{fix_h}: {w['fix']}")
            p.paragraph_format.left_indent = Inches(0.25)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _bullet(d, text: str) -> None:
    p = d.add_paragraph(style="List Bullet")
    _fill_runs(p, text)


def _fill_runs(p, text: str) -> None:
    from docx.enum.text import WD_COLOR_INDEX

    pos = 0
    for m in HILITE.finditer(text):
        p.add_run(text[pos:m.start()])
        r = p.add_run(m.group(0))
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    p.add_run(text[pos:])


# ---------------------------------------------------------------------------- professional layer
ADVICE_H = {"en": "Advisor's notes", "fr": "Conseils", "ar": "ملاحظات المستشار"}
UNMET = {"en": "Still to cover: {x}", "fr": "À couvrir : {x}", "ar": "للاستكمال: {x}"}
POLISH_NOTE = {"en": "Wording edited from your answers. Read every sentence against your own facts: anything marked 'to be confirmed' or 'NEEDED' was not in your answers.",
               "fr": "Rédaction reprise à partir de vos réponses. Relisez chaque phrase avec vos propres faits : tout ce qui est marqué « à confirmer » ou « NEEDED » ne figurait pas dans vos réponses.",
               "ar": "صياغة معدلة انطلاقا من إجاباتكم. راجعوا كل جملة مع معطياتكم: كل ما وُسم بـ«يؤكده الباعث» أو NEEDED لم يرد في إجاباتكم."}
REMOVED_NOTE = {"en": " {n} statement(s) the inputs did not support were replaced by questions.",
                "fr": " {n} affirmation(s) non étayée(s) par vos réponses ont été remplacée(s) par des questions.",
                "ar": " استُبدلت {n} جملة غير مدعومة بأسئلة."}

# facts every section sees, with an unambiguous label: the model must not mix one amount up with another
ANCHORS = {
    "project_name": "project_name",
    "stage": "stage",
    "legal_form": "legal_form",
    "headcount": "headcount_at_start",
    "total_investment": "total_investment_TND (everything needed to start, working capital included)",
    "funding_need_amount": "external_funding_sought_TND (the ask)",
    "funding_instrument": "main_funding_instrument",
    "own_funds": "own_funds_TND",
    "loan_amount": "loan_TND",
    "grant_amount": "grant_TND",
}


def advise(composed: dict, review, guidance_hits: list[dict], lang: str, critic_notes: list[str] | None = None) -> None:
    """Append advisor notes: critic demands first, then uncovered criteria, then the matching good practice. Works without a model."""
    from app.templates import label as _label

    notes: list[str] = list(critic_notes or [])
    if review is not None:
        for c in review.criteria:
            if not c.met:
                text = _label(c.text, lang) if isinstance(c.text, dict) else str(c.text)
                if "break-even" in text.lower() or "point mort" in text.lower():
                    continue  # never push the founder to state a break-even the model may not compute
                notes.append(UNMET.get(lang, UNMET["en"]).format(x=text))
    for g in guidance_hits[:1]:
        if "break-even" in f"{g['title']} {g['text']}".lower() or "point mort" in f"{g['title']} {g['text']}".lower():
            continue  # same rule as the criteria: never push a break-even the model may not compute
        notes.append(f"{g['title']}. {g['text']}")
    if notes:
        composed["blocks"].append(("h", ADVICE_H.get(lang, ADVICE_H["en"])))
        composed["blocks"].append(("ul", notes[:5]))
        composed["advice"] = notes[:5]


def _add_needed(composed: dict, items: list[str], lang: str) -> None:
    """Questions the writer could not answer from the facts join the section's 'To complete' list and the open items."""
    new = [i for i in items if i and not any(i in m for m in composed["missing"])]
    if not new:
        return
    entries = [f"[NEEDED: {i}]" for i in new]
    composed["missing"].extend(entries)
    blocks = composed["blocks"]
    todo = ("h", _t("todo", lang))
    for idx in range(len(blocks) - 1, 0, -1):
        if blocks[idx][0] == "ul" and blocks[idx - 1] == todo:
            blocks[idx][1].extend(entries)
            return
    blocks.append(todo)
    blocks.append(("ul", entries))


def polish(composed: dict, sec: dict, values: dict, derived: dict | None, guidance_hits: list[dict], lang: str,
           llm, masker, trace) -> bool:
    """Replace the founder's raw text blocks with plan prose. Every figure and every claim is checked against the founder's facts."""
    import json

    from app import prompts
    from app.agents import writer as W
    from app.config import model_for
    from app.llm import LLMUnavailable

    raw_ids = [q for q in sec["questions"] if q in values and q in BANK_BY_ID and BANK_BY_ID[q].type in ("text", "longtext")]
    if not llm.available or not raw_ids:
        return False
    cur = str(values.get("currency") or "TND")
    # The model sees every answer it is asked to rewrite, the founders included (the masker hides names, e-mails and phones),
    # plus labelled anchors so that "total investment" and "external funding sought" are never confused.
    facts = {"_context": {lab.replace("_TND", "_" + cur): values[k] for k, lab in ANCHORS.items() if k in values}}
    facts.update({q: values[q] for q in sec["questions"] if q in values})
    allowed = W.allowed_numbers(values, derived or {})
    evidence = claim_guard.build_evidence(values)
    marker = MARKERS.get(lang, MARKERS["en"])
    p = prompts.load("plan_writer")
    adv = "\n".join(f"- {g['title']}: {g['text']}" for g in guidance_hits[:2])
    feedback = ""
    for attempt in (0, 1):
        user = prompts.render(p["text"], section=composed["title"], language={"fr": "French", "ar": "Arabic", "en": "English"}[lang],
                              guidance=adv, facts=masker.mask(json.dumps(facts, ensure_ascii=False)), feedback=feedback,
                              currency=cur, marker=marker)
        with trace.step("plan_writer.llm" if attempt == 0 else "plan_writer.llm.retry", model=model_for("llm_main"), prompt_id=prompts.ref(p)):
            try:
                out = llm.chat_json(model=model_for("llm_main"), system="You write business-plan sections. Reply with JSON only.",
                                    user=user, schema=p.get("output_schema"), max_tokens=900)
            except LLMUnavailable as exc:
                trace.note(f"plan_writer skipped: {exc}")
                return False
        text = masker.unmask(str(out.get("text", ""))).strip()
        if not text:
            return False
        text = _UNKNOWN_TOKEN.sub("[NEEDED: confirm the name used here]", text)
        bad = W.unflagged_count(claim_guard.protect(text)[0], allowed)
        if bad and attempt == 0:
            feedback = f"\nYour previous draft contained {bad} figure(s) that are not in the facts. Remove them."
            continue
        text, removed = claim_guard.sanitize(text, evidence, currency=cur, marker=marker, lang=lang)
        protected, spans = claim_guard.protect(text)
        protected, _ = W.guard(protected, allowed, lang)
        text = claim_guard.restore(protected, spans)
        # only the mechanical [NEEDED] markers the writer left in the text; never the model's own list
        needed = claim_guard.needed_items(text)
        answered_labels = [{q, _qlabel(q, lang).lower()} for q in values if q in BANK_BY_ID]
        needed = [n for n in needed if not any(any(x and x in n.lower() for x in pair) for pair in answered_labels)]
        raw_labels = {_qlabel(q, lang) for q in raw_ids}
        kept, inserted = [], False
        i = 0
        blocks = composed["blocks"]
        while i < len(blocks):
            b = blocks[i]
            if b[0] == "h" and b[1] in raw_labels and i + 1 < len(blocks) and blocks[i + 1][0] == "p":
                if not inserted:
                    kept.append(("p", text))
                    inserted = True
                i += 2
                continue
            kept.append(b)
            i += 1
        composed["blocks"] = kept
        composed["polished"] = True
        composed["claims_removed"] = [dict(r) for r in removed]
        _add_needed(composed, needed, lang)
        return True
    return False
