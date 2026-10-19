"""Export the review and drafts as a Word document (artifact kind docx)."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt

from app.agents.writer import ANY_MARKER

T = {
    "en": {"title": "Business plan review", "summary": "Summary", "completeness": "Completeness by section", "signals": "Points to review",
           "fin": "Financial check", "drafts": "Drafted sections", "assump": "Assumptions", "ctx": "What this review is based on",
           "note": "This report lists points to review. It does not give legal, tax or lending advice. Figures marked as to be confirmed are not in your inputs.",
           "sec": "Section", "score": "Score", "status": "Status", "sev": "Level", "msg": "Point", "field": "Field", "src": "Source", "val": "Value"},
    "fr": {"title": "Revue du business plan", "summary": "Synthèse", "completeness": "Complétude par section", "signals": "Points à vérifier",
           "fin": "Contrôle financier", "drafts": "Sections rédigées", "assump": "Hypothèses", "ctx": "Base de cette revue",
           "note": "Ce rapport liste des points à vérifier. Il ne constitue pas un conseil juridique, fiscal ou bancaire. Les chiffres marqués à confirmer ne figurent pas dans vos données.",
           "sec": "Section", "score": "Score", "status": "Statut", "sev": "Niveau", "msg": "Point", "field": "Champ", "src": "Source", "val": "Valeur"},
}
T["ar"] = T["en"]


def _para_with_markers(doc, text: str) -> None:
    for line in text.split("\n"):
        p = doc.add_paragraph()
        pos = 0
        for m in ANY_MARKER.finditer(line):
            p.add_run(line[pos:m.start()])
            r = p.add_run(m.group(0))
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            pos = m.end()
        p.add_run(line[pos:])


def _table(doc, header: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)


def build_docx(report: dict, chart_png: bytes | None, lang: str) -> bytes:
    L = T.get(lang, T["en"])
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading(L["title"], 0)
    d.add_paragraph(report["summary_text"])
    d.add_paragraph(L["note"]).italic = True

    a = report.get("assessment")
    if a:
        d.add_heading({"fr": "Évaluation de crédibilité", "ar": "تقييم الموثوقية"}.get(lang, "Credibility assessment"), 1)
        d.add_paragraph(f"{a['label'].get(lang) or a['label']['en']} ({a['score']}/100). {a['rationale']}")
        rows = [[f["group"], f["why"], (f.get("quote") or "")[:160]] for f in a["flags"]]
        if rows:
            _table(d, ["type", "why", "quote"], rows)
        for st in a["strengths"]:
            d.add_paragraph(st["why"], style="List Bullet")

    bn = report.get("benchmarks") or {}
    if bn.get("sector") or bn.get("notes"):
        d.add_heading("Reference data", 1)
        if bn.get("sector"):
            b = bn["sector"]["benchmark"]
            d.add_paragraph(f"Sector matched: {bn['sector']['matched_industry']}. Reference after-tax operating margin {b['after_tax_operating_margin']:.0%} "
                            f"({b['region']} firms, {b['as_of']}, Damodaran). US data: a yardstick, not a Tunisian norm.")
        for n in bn.get("notes", []):
            d.add_paragraph(n, style="List Bullet")

    d.add_heading(L["ctx"], 1)
    ctx = report["context"]
    rows = [[k, str(v)[:120], ctx["provenance"].get(k, "")] for k, v in ctx["values"].items() if not isinstance(v, (list, dict))]
    if rows:
        _table(d, [L["field"], L["val"], L["src"]], rows)
    for c in ctx["conflicts"]:
        d.add_paragraph(f"Conflict on {c['field']}: " + ", ".join(f"{s} = {v}" for s, v in c["values"].items()) + f". Used: {c['chosen_source']}.")

    d.add_heading(L["completeness"], 1)
    _table(d, [L["sec"], L["score"], L["status"]], [[s["title"], s["score"], s["status"]] for s in report["completeness"]["sections"]])

    d.add_heading(L["signals"], 1)
    if report["signals"]:
        _table(d, [L["sev"], L["msg"]], [[s["severity"], s["message"]] for s in report["signals"]])
    else:
        d.add_paragraph("-")

    fin = report.get("financials")
    if fin and fin.get("years"):
        d.add_heading(L["fin"], 1)
        keys = [("revenue", "Revenue"), ("ebitda", "EBITDA"), ("net_income", "Net result"), ("cash_end", "Year-end cash")]
        _table(d, [""] + fin["years"], [[lab] + ["-" if v is None else f"{v:,.0f}".replace(",", " ") for v in fin[k]] for k, lab in keys])
        if chart_png:
            d.add_picture(io.BytesIO(chart_png), width=Inches(6.2))
        if fin.get("assumptions"):
            d.add_heading(L["assump"], 2)
            for a in fin["assumptions"]:
                d.add_paragraph(a, style="List Bullet")

    if report["drafts"]:
        d.add_heading(L["drafts"], 1)
        for dr in report["drafts"]:
            d.add_heading(dr["title"], 2)
            _para_with_markers(d, dr["text"])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
