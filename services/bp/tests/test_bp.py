"""Contract, questionnaire, merge, checker, writer guard, pipeline and the mock."""

import io
import json

import httpx
import pytest
from fastapi.testclient import TestClient

from app import storage
from app.agents.writer import allowed_numbers, guard, unflagged_count
from app.contract import Request, Response
from app.main import app
from app.profile import merge_context
from app.questionnaire import clean_answers, parse_number
from eval.run_eval import financial as eval_financial, merge as eval_merge, writer as eval_writer

client = TestClient(app)

PLAN = """RÉSUMÉ DU PROJET
Application pour cliniques, financement de 120 000 TND.

Promoteurs et équipe
Le fondateur a huit ans d'expérience, rôle de gérant.

Plan de financement
Emplois   Montant   Ressources   Montant
Investissements matériels   70 000   Fonds propres   40 000
Fonds de roulement   50 000   Crédit bancaire   100 000
Total emplois   120 000   Total ressources   140 000

Compte de résultat
Poste   2026   2027   2028
Chiffre d'affaires   100 000   180 000   260 000
Achats consommés   30 000   50 000   70 000
Charges de personnel   60 000   70 000   80 000
Résultat net   90 000   20 000   55 000
"""

ANSWERS = {"problem": "Les cliniques perdent du temps.", "solution": "Un logiciel de rendez-vous.", "horizon_years": "3",
           "revenue_by_year": "100000, 180000, 260000", "stage": "extension"}


def test_endpoints_exist():
    assert client.get("/health").json()["status"] == "ok"
    meta = client.get("/v1/meta").json()
    assert meta["module"] == "bp" and "review" in meta["modes"]


def test_abstain_without_input():
    r = client.post("/v1/run", json={"request_id": "1", "workspace": "bp"})
    body = Response(**r.json())
    assert body.status == "abstain"


def test_bad_mode_is_bad_input():
    r = client.post("/v1/run", json={"request_id": "1", "options": {"mode": "x", "answers": {"sector": "a"}}})
    assert r.status_code == 400 and r.json()["error"]["code"] == "BAD_INPUT"


def test_unknown_file_ref_and_traversal():
    r = client.post("/v1/run", json={"request_id": "1", "file_ref": "../../etc/passwd"})
    assert r.status_code == 400
    r = client.post("/v1/run", json={"request_id": "1", "file_ref": "a" * 64})
    assert r.status_code == 400


def test_legacy_office_refused():
    ref = storage.put(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 50)
    r = client.post("/v1/run", json={"request_id": "1", "file_ref": ref})
    assert r.status_code == 415 and r.json()["error"]["code"] == "UNSUPPORTED_FILE"


def test_number_parsing():
    assert parse_number("50 000") == 50000 and parse_number("1,5") == 1.5 and parse_number("12k") == 12000


def test_answers_validation():
    c = clean_answers({"loan_rate_pct": "120", "legal_form": "sarl", "nonsense": 1})
    assert "loan_rate_pct" in c.errors and c.values["legal_form"] == "SARL" and c.unknown == ["nonsense"]


def test_merge_precedence_and_conflict():
    c = merge_context({"stage": "creation", "sector": "software"}, {"stage": "extension"}, {"legal_form": "SARL"})
    assert c.values["stage"] == "extension" and c.provenance["stage"] == "questionnaire"
    assert c.provenance["sector"] == "orchestrator" and c.provenance["legal_form"] == "document"
    assert c.conflicts and c.conflicts[0]["field"] == "stage"


def test_questionnaire_marks_prefilled():
    r = client.post("/v1/questionnaire", json={"lang": "fr", "profile": {"stage": "creation", "sector": "bois"}, "answers": {"problem": "x"}})
    q = {x["id"]: x for x in r.json()["questions"]}
    assert q["sector"]["status"] == "prefilled" and q["problem"]["status"] == "answered" and q["pricing"]["status"] == "missing"


def test_questionnaire_only_run_has_artifacts_and_report():
    r = client.post("/v1/run", json={"request_id": "q", "lang": "en", "profile": {"stage": "creation", "sector": "software"},
                                    "options": {"answers": {**ANSWERS, "stage": "creation", "total_investment": "80 000", "own_funds": "30000",
                                                           "loan_amount": "50000", "loan_rate_pct": "9", "loan_years": "5"}}})
    assert r.status_code == 200
    body = Response(**r.json())
    kinds = {a.kind for a in body.artifacts}
    assert kinds == {"png", "docx", "json"}
    rep = json.loads(storage.read([a for a in body.artifacts if a.kind == "json"][0].file_ref))
    assert set(rep["input"]["methods"]) == {"questionnaire", "orchestrator_profile"}
    assert rep["context"]["provenance"]["sector"] == "orchestrator"
    assert all(s.evidence is None or s.severity in ("info", "review", "high") for s in body.signals)
    assert all("review this" in s.message.lower() for s in body.signals)


def test_upload_plus_questionnaire_flags_everything():
    ref = storage.put(PLAN.encode())
    r = client.post("/v1/run", json={"request_id": "u", "lang": "fr", "file_ref": ref,
                                    "profile": {"stage": "creation"}, "options": {"answers": {**ANSWERS, "revenue_by_year": "100000, 150000, 260000"}}})
    body = Response(**r.json())
    kinds = {s.kind for s in body.signals}
    assert "sum_mismatch" in kinds          # stated net result 90 000 is not revenue minus costs; totals differ
    assert "sources_neq_uses" in kinds      # 140 000 sources versus 120 000 uses
    assert "source_conflict" in kinds       # 180 000 in the plan, 150 000 in the questionnaire
    assert "profile_conflict" in kinds      # creation (orchestrator) vs extension (questionnaire)
    assert body.status == "ok"
    rep = json.loads(storage.read([a for a in body.artifacts if a.label == "bp_report"][0].file_ref))
    assert set(rep["input"]["methods"]) == {"upload", "questionnaire", "orchestrator_profile"}


def test_draft_section_marks_unknown_figures():
    r = client.post("/v1/run", json={"request_id": "d", "lang": "en", "options": {"mode": "draft_section", "section_id": "financing_plan", "answers": {"own_funds": "30000"}}})
    body = Response(**r.json())
    assert "[to be confirmed by the founder]" in body.answer
    assert "30 000" in body.answer


def test_draft_section_unknown_id():
    r = client.post("/v1/run", json={"request_id": "d", "options": {"mode": "draft_section", "section_id": "nope", "answers": {"sector": "a"}}})
    assert r.status_code == 400


def test_handoff_for_statute_and_aid_questions():
    r = client.post("/v1/run", json={"request_id": "h", "question": "Que dit l'article 92 du code des sociétés ?", "options": {"answers": {"sector": "a"}}})
    assert r.json()["handoff"] == {"to": "law", "question": "Que dit l'article 92 du code des sociétés ?"}
    r = client.post("/v1/run", json={"request_id": "h", "question": "Quel programme BFPME convient ?", "options": {"answers": {"sector": "a"}}})
    assert r.json()["handoff"]["to"] == "aid"


def test_masking_hides_identifiers_from_the_model():
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request.content.decode())
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps({"sections": []})}}]})

    from app.agents.common import Trace
    from app.agents import reviewer
    from app.llm import LLMClient
    from app.masking import Masker
    from app.templates import load_template

    llm = LLMClient(api_key="k", transport=httpx.MockTransport(handler), cache=False, max_retries=0)
    reviewer.review_all(load_template("bfpme_creation"), {"promoters_team": {"text": "Ahmed Ben Salah, MF 1234567/A/M/000, tél +216 98 123 456"}},
                        lang="en", llm=llm, masker=Masker(names=["Ahmed Ben Salah"]), trace=Trace())
    assert seen and "Ahmed" not in seen[0] and "1234567" not in seen[0] and "98 123" not in seen[0]


def test_reviewer_discards_fabricated_quotes():
    def handler(request):
        crit = [{"id": "c_exp", "met": True, "evidence_quote": "ten years at Google"}, {"id": "c_roles", "met": True, "evidence_quote": "role de gérant"}]
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps({"sections": [{"section_id": "promoters_team", "criteria": crit, "critique": "ok"}]})}}]})

    from app.agents.common import Trace
    from app.agents import reviewer
    from app.llm import LLMClient
    from app.masking import Masker
    from app.templates import load_template

    llm = LLMClient(api_key="k", transport=httpx.MockTransport(handler), cache=False, max_retries=0)
    rv = reviewer.review_all(load_template("bfpme_creation"), {"promoters_team": {"text": "Le fondateur, role de gérant."}}, lang="en", llm=llm, masker=Masker(), trace=Trace())
    team = next(r for r in rv if r.section_id == "promoters_team")
    got = {c.id: c for c in team.criteria}
    assert not got["c_exp"].met and got["c_exp"].note and got["c_roles"].met and team.method == "model"


def test_llm_failure_degrades_not_crashes():
    def handler(request):
        return httpx.Response(500)

    from app.llm import LLMClient, LLMUnavailable
    llm = LLMClient(api_key="k", transport=httpx.MockTransport(handler), cache=False, max_retries=0)
    with pytest.raises(LLMUnavailable):
        llm.chat_json(model="m", system="s", user="u")


def test_retry_after_on_429():
    calls = {"n": 0}

    def handler(request):
        calls["n"] += 1
        if calls["n"] == 1:
            return httpx.Response(429, headers={"Retry-After": "0"})
        return httpx.Response(200, json={"choices": [{"message": {"content": "{\"a\": 1}"}}], "usage": {"prompt_tokens": 3, "completion_tokens": 2}})

    from app.llm import LLMClient
    llm = LLMClient(api_key="k", transport=httpx.MockTransport(handler), cache=False, max_retries=2)
    assert llm.chat_json(model="m", system="s", user="u") == {"a": 1}
    assert llm.usage.rate_limited == 1 and llm.usage.prompt_tokens == 3


def test_guard_marks_untraceable_numbers():
    al = allowed_numbers({"a": 100000})
    text, flagged = guard("CA de 100 000 TND puis 300 000 TND.", al, "fr")
    assert len(flagged) == 1 and unflagged_count(text, al) == 0 and unflagged_count("300 000 TND", al) == 1


def test_eval_suites_meet_floor():
    f = eval_financial(7, None)
    assert f["recall_overall"] >= 0.9 and f["false_positive_rate"] <= 0.1
    assert eval_writer(7, 6)["unflagged_after_guard"] == 0
    assert eval_merge(7, None)["accuracy"] == 1.0


def test_docx_artifact_opens():
    import docx
    r = client.post("/v1/run", json={"request_id": "x", "options": {"answers": ANSWERS}})
    body = Response(**r.json())
    ref = next(a.file_ref for a in body.artifacts if a.kind == "docx")
    d = docx.Document(io.BytesIO(storage.read(ref)))
    assert len(d.paragraphs) > 3


def test_mock_matches_contract(uploads):
    from mock.main import app as mock_app
    m = TestClient(mock_app)
    assert Response(**m.post("/v1/run", json={"request_id": "m"}).json()).status == "abstain"
    ok = Response(**m.post("/v1/run", json={"request_id": "m", "options": {"answers": {"sector": "a"}}}).json())
    assert ok.status == "ok" and storage.read(ok.artifacts[0].file_ref)


def _run_file(data: bytes):
    ref = storage.put(data)
    r = client.post("/v1/run", json={"request_id": "f", "lang": "fr", "file_ref": ref})
    assert r.status_code == 200, r.text
    return Response(**r.json())


def test_docx_upload_with_table():
    import docx
    d = docx.Document()
    d.add_heading("Plan de financement", 1)
    d.add_paragraph("Le financement combine fonds propres et crédit.")
    t = d.add_table(rows=4, cols=2)
    for i, (a, b) in enumerate([("Emplois", "Montant"), ("Investissements matériels", "70 000"), ("Total emplois", "70 000"), ("Fonds propres", "50 000")]):
        t.rows[i].cells[0].text, t.rows[i].cells[1].text = a, b
    t2 = d.add_table(rows=3, cols=3)
    for i, row in enumerate([("Poste", "2026", "2027"), ("Chiffre d'affaires", "100000", "90000"), ("Achats consommés", "20000", "20000")]):
        for j, v in enumerate(row):
            t2.rows[i].cells[j].text = v
    buf = io.BytesIO(); d.save(buf)
    body = _run_file(buf.getvalue())
    assert body.status == "ok"
    rep = json.loads(storage.read(next(a.file_ref for a in body.artifacts if a.label == "bp_report")))
    assert rep["input"]["document"]["kind"] == "docx" and rep["financials"]["revenue"][:2] == [100000.0, 90000.0]


def test_xlsx_upload():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Compte de résultat"
    ws.append(["Poste", 2026, 2027, 2028]); ws.append(["Chiffre d'affaires", 100000, 180000, 260000]); ws.append(["Charges de personnel", 60000, 70000, 80000])
    buf = io.BytesIO(); wb.save(buf)
    body = _run_file(buf.getvalue())
    rep = json.loads(storage.read(next(a.file_ref for a in body.artifacts if a.label == "bp_report")))
    assert rep["financials"]["revenue"] == [100000.0, 180000.0, 260000.0]


def test_french_presentation_maps_project_and_executive():
    from pathlib import Path

    from app.agents.common import Trace
    from app.agents import parser as parser_agent
    from app.intake import read_document
    from app.llm import LLMClient
    from app.masking import Masker
    from app.templates import load_template

    root = Path(__file__).resolve().parents[3] / "plan_dataset" / "raw" / "invented_tn" / "g1_huilerie_zeitouna.docx"
    if not root.exists():
        return
    doc = read_document(root.read_bytes())
    tpl = load_template("bfpme_creation")
    plan = parser_agent.parse(doc, tpl, llm=LLMClient(api_key="", cache=False), masker=Masker(), lang="fr", trace=Trace(), mode="keywords")
    assert plan.sections.get("project_description", {}).get("text", "").strip()
    assert plan.sections.get("executive_summary", {}).get("text", "").strip()


def test_pdf_upload_and_scan_flag():
    import pymupdf
    pdf = pymupdf.open(); pg = pdf.new_page(); pg.insert_text((72, 72), "Etude de marche", fontsize=18)
    pg.insert_text((72, 110), "La clientele cible est composee de cliniques. Concurrents: deux editeurs.", fontsize=10)
    pdf.new_page()  # blank page, no text layer
    body = _run_file(pdf.tobytes())
    rep = json.loads(storage.read(next(a.file_ref for a in body.artifacts if a.label == "bp_report")))
    assert rep["input"]["document"]["scanned_pages"] == [2]
    assert any(s["section_id"] == "market_analysis" and s["status"] != "missing" for s in rep["completeness"]["sections"])


# ---- chat questionnaire and credibility (added with the pitch-case fix) --------------
def test_chat_flow_and_skip():
    from app import interview as iv

    r = iv.turn(lang="en", profile={"sector": "SaaS"}, state=None, message=None)
    assert r["asking"] == "project_name" and any("already know" in m for m in r["messages"])
    st = r["state"]
    r = iv.turn(lang="en", profile=None, state=st, message="Medilink")
    assert r["state"]["answers"]["project_name"] == "Medilink"
    asked = r["asking"]
    r = iv.turn(lang="en", profile=None, state=r["state"], message="skip")
    assert asked in r["state"]["skipped"] and r["asking"] != asked


def test_chat_numbers_and_years():
    from app import interview as iv

    v, err = iv.parse_reply("total_investment", "about 200k TND", {})
    assert v == 200000 and err is None
    v, err = iv.parse_reply("revenue_by_year", "10k, 20k", {"horizon_years": 3})
    assert v is None and err == "years"
    v, err = iv.parse_reply("revenue_by_year", "10k, 20k, 40k", {})
    assert v == [10000, 20000, 40000]


def test_chat_reaches_review_offer():
    from app import interview as iv

    st = iv.new_state()
    r = iv.turn(lang="fr", profile=None, state=st, message=None)
    for _ in range(40):
        if r["phase"] != "core":
            break
        r = iv.turn(lang="fr", profile=None, state=r["state"], message=None, action="skip")
    assert r["phase"] == "offer" and r["ready_to_review"]


def test_pitch_cases_match_answer_key():
    from eval.run_pitch_eval import main

    assert main(verbose=False) >= 0.9


# ---- reference data from the pack ---------------------------------------------------
def test_sector_mapping_and_benchmark():
    from app import knowledge as kb

    assert kb.industry_for("SaaS for clinics") == "Software (System & Application)"
    assert kb.industry_for("huile d'olive") == "Farming/Agriculture"
    assert kb.industry_for("zzz") is None
    bm = kb.industry_benchmark("Food Processing")
    assert bm["region"] == "US" and 0 < bm["after_tax_operating_margin"] < 0.3


def test_benchmark_checks_flag_and_label_source():
    from app import benchmark_check as bc
    from app import financial as fin

    m = fin.FinancialModel(years=["Y1", "Y2", "Y3"], revenue=[100_000, 200_000, 300_000], cogs=[20_000, 30_000, 40_000], payroll=[10_000, 12_000, 14_000],
                           opex=[5_000, 6_000, 7_000], loan_rate_pct=2.0, tax_rate_pct=35.0, headcount=2, avg_gross_salary=500, origin="questionnaire")
    d = fin.derive(m)
    sig, rep = bc.run(m, d, {"sector": "boulangerie"})
    kinds = {s.kind for s in sig}
    assert {"margin_above_sector", "loan_rate_low", "tax_rate_unusual"} <= kinds
    assert "US" in next(s.message for s in sig if s.kind == "margin_above_sector")
    assert rep["sector"]["matched_industry"] == "Food Processing"
    assert any("secondary" in n for n in rep["notes"])


def test_benchmark_checks_quiet_on_plausible_plan():
    from app import benchmark_check as bc
    from app import financial as fin

    m = fin.FinancialModel(years=["Y1"], revenue=[100_000], cogs=[50_000], payroll=[30_000], opex=[10_000], loan_rate_pct=9.5, tax_rate_pct=20.0, origin="questionnaire")
    sig, _ = bc.run(m, fin.derive(m), {"sector": "boulangerie"})
    assert not sig


def test_build_knowledge_and_search(tmp_path, monkeypatch):
    import csv
    import subprocess
    import sys

    from app import knowledge as kb

    pack = tmp_path / "pack"
    (pack / "catalog").mkdir(parents=True)
    (pack / "t").mkdir()
    (pack / "t" / "S1_code.txt").write_text("\n".join(f"Article {i}\nLe taux de l'impôt sur les sociétés est de vingt pour cent. " + "texte " * 40 for i in range(1, 8)), encoding="utf-8")
    with (pack / "catalog" / "file_index.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["pack_path", "source_id", "title", "kind", "topic", "language", "pages", "text_chars", "has_text_extract", "source_url", "last_modified", "sha256", "use"])
        w.writerow(["t/S1_code.pdf", "S1", "Code IS", "pdf", "tax", "fr", "1", "1", "True", "http://x", "", "", "rag"])
    out = tmp_path / "k.sqlite"
    subprocess.run([sys.executable, "scripts/build_knowledge.py", str(pack), "--out", str(out)], check=True, capture_output=True)
    monkeypatch.setattr(kb, "DATA", out.parent)
    (out.parent / "knowledge.sqlite").write_bytes(out.read_bytes())
    kb._kb.cache_clear()
    hits = kb.search("impôt sociétés taux", k=1)
    kb._kb.cache_clear()
    assert hits and hits[0]["doc_id"] == "S1" and "Article" in hits[0]["title"] and hits[0]["url"] == "http://x"


def test_chat_understands_natural_answers():
    from app import interview as iv

    assert iv.parse_reply("stage", "we are creating it", {})[0] == "creation"
    assert iv.parse_reply("stage", "already running for 3 years", {})[0] == "extension"
    assert iv.parse_reply("funding_instrument", "a bank loan", {})[0] == "loan"
    for t in ("100000, 250000, 400000", "100k, 250k, 400k", "100 000; 250 000; 400 000", "year 1: 100000, year 2: 250000, year 3: 400000"):
        assert iv.parse_reply("revenue_by_year", t, {})[0] == [100000, 250000, 400000], t
    assert iv.parse_reply("revenue_by_year", "100000, 250000", {})[1] == "years"


def test_generate_plan_artifact_from_answers():
    from app import storage
    from app.contract import Request
    from app.pipeline import run
    import docx, io
    ans = {"project_name": "Atelier Test", "sector": "artisanat", "problem": "Les artisans vendent mal en ligne",
           "solution": "Une marketplace pour artisans", "total_investment": 100000, "horizon_years": 3,
           "revenue_by_year": [100000, 250000, 400000]}
    r = run(Request(request_id="t", lang="fr", options={"answers": ans, "generate_plan": True}), None)
    labels = [a.label for a in r.artifacts]
    assert "business_plan" in labels and "bp_review" in labels
    art = [a for a in r.artifacts if a.label == "business_plan"][0]
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(storage.read(art.file_ref))).paragraphs)
    assert "Atelier Test" in text and "Les artisans vendent mal en ligne" in text
    assert "à confirmer" in text  # missing items are marked, not invented
    assert "point mort est atteint" not in text  # no cost lines: no break-even claim


class _StubLLM:
    available = True
    deadline = None

    class usage:  # noqa: N801
        calls = 1
        cache_hits = 0

        @staticmethod
        def as_dict():
            return {}

    def __init__(self, text):
        self.text = text

    def chat_json(self, **kw):
        return {"text": self.text}


def test_professional_layer_rewrites_and_guards_numbers():
    from app import storage
    from app.contract import Request
    from app.pipeline import run
    import docx, io
    ans = {"project_name": "Atelier Test", "sector": "artisanat", "problem": "les artisans vendent mal sur internet",
           "solution": "un site pour vendre leurs produits", "total_investment": 100000}
    stub = _StubLLM("Les artisans tunisiens peinent à vendre en ligne. Atelier Test propose une place de marché dédiée, "
                    "avec 7777 clients attendus.")
    r = run(Request(request_id="t", lang="fr", options={"answers": ans, "generate_plan": True, "llm": True}), stub)
    art = [a for a in r.artifacts if a.label == "business_plan"][0]
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(storage.read(art.file_ref))).paragraphs)
    assert "place de marché dédiée" in text           # professional wording used
    assert "7777 [à confirmer par le fondateur]" in text  # invented figure is marked, never presented as fact
    assert "Conseils" in text                         # advisor notes present


def test_model_name_can_be_overridden_by_env(monkeypatch):
    from app.config import model_for
    monkeypatch.setenv("LLM_MODEL_MAIN", "qwen2.5:7b")
    assert model_for("llm_main") == "qwen2.5:7b"
    monkeypatch.delenv("LLM_MODEL_MAIN")
    assert model_for("llm_main") != "qwen2.5:7b"


def test_rerank_prefers_prose_over_account_list_fragments():
    from app.knowledge import rerank
    terms = ["amortissement", "immobilisations"]
    fragment = "142 Résultat\\nde l'exercice\\n15 Régularisations\\n151 Somme\\nimmobilisations"
    prose = ("Les immobilisations corporelles sont amorties selon le mode adopté. L'amortissement est calculé "
             "sur la durée d'utilité estimée de chaque immobilisation et indiqué dans les notes.")
    assert rerank(5.0, prose, terms) > rerank(5.0, fragment, terms)


def test_llm_steps_gate(monkeypatch):
    from app.config import settings
    monkeypatch.setenv("BP_LLM_STEPS", "plan_writer,drafts")
    assert settings.llm_step("plan_writer") and not settings.llm_step("reviewer")
    monkeypatch.delenv("BP_LLM_STEPS")
    assert settings.llm_step("reviewer")


def test_unfilled_template_is_detected_and_not_scored_as_plan():
    from app.template_detect import scan
    tpl = ("Executive Summary\nHelp tip\nStart writing here..\nMarket opportunity\nHelp tip\nStart writing here..\n"
           "Start writing here..\nStart writing here..\nJohn Doe CEO john.doe@example.com\n"
           "[John Doe] brings over [XX] years of experience. [He/She] holds a [Type of Degree] from [University/Institution Name].\n"
           "Business Plan [YEAR]\n")
    assert scan(tpl)["unfilled"]
    real = "Notre huilerie conditionne 40 tonnes d'huile d'olive par an à Sfax. Investissement total 180 000 TND. " * 5
    assert not scan(real)["unfilled"]


def test_prose_revenue_and_funding_ask():
    from app.agents.parser import prose_funding_ask, prose_revenue
    assert prose_revenue("Financials\nYear 1 revenue 400k, Year 2 2M, Year 3 6M. We need around 300k.") == (["Y1", "Y2", "Y3"], [400000.0, 2000000.0, 6000000.0])
    assert prose_revenue("Year 1: 150 accounts, ARR 520k MAD, burn 1.6M MAD. Year 2: 700 accounts, ARR 2.9M MAD.") == (["Y1", "Y2"], [520000.0, 2900000.0])
    assert prose_revenue("Profit of 100M MAD in the first year.") is None
    assert prose_funding_ask("Seeking 2.5M MAD for 24 months") == 2500000.0


# ---- final institutional writer (final_writer@v3 batched, v2 per section) --------------
def test_final_writer_agent_writes_six_guarded_sections():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw

    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode())
        seen.append(body)
        sid = next(s for s in fw.SECTION_IDS if f"Write ONLY section {fw.SECTION_TITLES[s]}" in body["messages"][1]["content"])
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
            {"id": sid, "title": fw.SECTION_TITLES[sid], "text": "The plan rests on 80 000 TND of own funds and targets 150 000 TND revenue."})}}],
            "usage": {"prompt_tokens": 5, "completion_tokens": 4}})

    out = asyncio.run(fw.run_final_writer_agent(prompts, '{"total_investment": "80 000"}', "plan text", "[]",
                                                transport=httpx.MockTransport(handler), mode="sequential"))
    assert [s["id"] for s in out["sections"]] == fw.SECTION_IDS
    assert all(s["status"] == "ok" for s in out["sections"]) and out["method"] == "model" and out["mode"] == "sequential"
    assert len(seen) == 6 and all(b["temperature"] == 0 and b["seed"] == 7 for b in seen)
    text = out["sections"][0]["text"]
    assert "80 000 TND" in text and "150 000 TND [to be confirmed by the founder]" in text


def test_final_writer_agent_batch_is_one_call():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw

    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode())
        seen.append(body)
        assert fw.BATCH_MARKER in body["messages"][1]["content"]
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
            {"sections": [{"id": sid, "title": fw.SECTION_TITLES[sid],
                           "text": "Uses 80 000 TND of own funds and targets 150 000 TND revenue."} for sid in fw.SECTION_IDS]})}}],
            "usage": {"prompt_tokens": 5, "completion_tokens": 4}})

    out = asyncio.run(fw.run_final_writer_agent(prompts, '{"total_investment": "80 000"}', "plan text", "[]",
                                                transport=httpx.MockTransport(handler)))
    assert out["mode"] == "batch" and out["method"] == "model" and out["usage"]["calls"] == 1
    assert len(seen) == 1 and seen[0]["temperature"] == 0 and seen[0]["seed"] == 7
    assert [s["id"] for s in out["sections"]] == fw.SECTION_IDS
    assert "150 000 TND [to be confirmed by the founder]" in out["sections"][5]["text"]


def test_final_writer_agent_batch_tops_up_missing_sections():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw

    users = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode())
        user = body["messages"][1]["content"]
        users.append(user)
        if fw.BATCH_MARKER in user:
            keep = fw.SECTION_IDS[:3]
            return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
                {"sections": [{"id": sid, "title": fw.SECTION_TITLES[sid], "text": "batch body."} for sid in keep]})}}]})
        sid = next(s for s in fw.SECTION_IDS if f"Write ONLY section {fw.SECTION_TITLES[s]}" in user)
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
            {"id": sid, "title": fw.SECTION_TITLES[sid], "text": "topped-up body."})}}]})

    out = asyncio.run(fw.run_final_writer_agent(prompts, "{}", "plan", "[]", transport=httpx.MockTransport(handler)))
    assert out["method"] == "model" and out["usage"]["calls"] == 4 and len(users) == 4
    assert all(s["status"] == "ok" for s in out["sections"])
    assert out["sections"][0]["text"] == "batch body." and out["sections"][3]["text"] == "topped-up body."


def test_final_writer_agent_retries_429():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw
    calls = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        calls["n"] += 1
        if calls["n"] == 1:
            return httpx.Response(429, headers={"Retry-After": "0"})
        payload = {"sections": [{"id": sid, "text": f"body {sid}."} for sid in fw.SECTION_IDS]}
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(payload)}}], "usage": {}})

    out = asyncio.run(
        fw.run_final_writer_agent(
            prompts, "{}", "plan", "[]", transport=httpx.MockTransport(handler), max_retries=2,
        )
    )
    assert out["method"] == "model" and calls["n"] == 2
    assert all(s["status"] == "ok" for s in out["sections"])


def test_final_writer_agent_degrades_not_crashes():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw
    calls = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        calls["n"] += 1
        return httpx.Response(500)

    out = asyncio.run(fw.run_final_writer_agent(prompts, "{}", "plan", "[]",
                                                transport=httpx.MockTransport(handler), max_retries=0))
    assert [s["id"] for s in out["sections"]] == fw.SECTION_IDS
    assert all(s["status"] == "error" for s in out["sections"]) and out["method"] == "unavailable"
    assert calls["n"] == 7  # one batch attempt, then one single attempt per section, no retry storm


def test_final_writer_agent_stops_when_budget_gone():
    import asyncio
    import time

    from app import prompts
    from app.agents import final_writer as fw
    calls = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        calls["n"] += 1
        return httpx.Response(200, json={"choices": [{"message": {"content": "{}"}}]})

    out = asyncio.run(fw.run_final_writer_agent(prompts, "{}", "plan", "[]", transport=httpx.MockTransport(handler),
                                                deadline=time.monotonic() - 1))
    assert calls["n"] == 0 and all(s["status"] == "error" for s in out["sections"])
    assert "time budget" in out["sections"][0]["error"]


def test_final_writer_agent_skipped_without_key():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw

    out = asyncio.run(fw.run_final_writer_agent(prompts, "{}", "plan", "[]"))
    assert out["method"] == "unavailable" and all(s["status"] == "skipped" for s in out["sections"])


def test_pipeline_generate_plan_runs_critic_and_final_writer(monkeypatch):
    from app import pipeline as pl
    from app.agents import final_writer as fw
    from app.llm import LLMClient

    monkeypatch.setenv("BP_LLM_STEPS", "critic,final_writer,plan_writer,drafts")
    monkeypatch.setenv("LLM_API_KEY", "k")

    def sync_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
            {"weaknesses": [{"dimension": "market_realism", "quote": None, "section_id": "market_analysis",
                             "severity": "medium", "why": "no TAM given", "fix": "add a TAM/SAM/SOM breakdown"}]})}}],
            "usage": {"prompt_tokens": 3, "completion_tokens": 2}})

    async def fake_post(client, body, remaining):
        user = body["messages"][1]["content"]
        if fw.BATCH_MARKER in user:
            return json.dumps({"sections": [{"id": sid, "title": fw.SECTION_TITLES[sid],
                                             "text": f"{fw.SECTION_TITLES[sid]}: revenue target 999 999 TND."} for sid in fw.SECTION_IDS]}), {"prompt_tokens": 5, "completion_tokens": 4}
        sid = next(s for s in fw.SECTION_IDS if f"Write ONLY section {fw.SECTION_TITLES[s]}" in user)
        return json.dumps({"id": sid, "title": fw.SECTION_TITLES[sid], "text": "topped-up body."}), {"prompt_tokens": 5, "completion_tokens": 4}

    monkeypatch.setattr(fw, "_post", fake_post)
    llm = LLMClient(api_key="k", transport=httpx.MockTransport(sync_handler), cache=False, max_retries=0)
    r = pl.run(Request(request_id="fp", lang="en", options={"answers": ANSWERS, "generate_plan": True}), llm)
    assert r.status == "ok"
    rep = json.loads(storage.read(next(a.file_ref for a in r.artifacts if a.label == "bp_report")))
    final_plan = rep["final_plan"]
    assert [s["id"] for s in final_plan["sections"]] == fw.SECTION_IDS
    assert all(s["status"] == "ok" for s in final_plan["sections"])
    assert final_plan["method"] == "model" and final_plan["mode"] == "batch" and final_plan["usage"]["calls"] == 1
    financial = final_plan["sections"][-1]
    assert "999 999 TND [to be confirmed by the founder]" in financial["text"]


def test_pipeline_rejects_unknown_final_writer_mode():
    from app import pipeline as pl
    from app.contract import ModuleError

    with pytest.raises(ModuleError):
        pl.run(Request(request_id="x", options={"answers": {"sector": "a"}, "final_writer_mode": "nope"}), None)


def test_final_writer_agent_seed_is_threaded():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw

    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(json.loads(request.content.decode()))
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
            {"sections": [{"id": sid, "title": fw.SECTION_TITLES[sid], "text": "body."} for sid in fw.SECTION_IDS]})}}]})

    asyncio.run(fw.run_final_writer_agent(prompts, "{}", "plan", "[]", transport=httpx.MockTransport(handler),
                                          seed=42))
    assert len(seen) == 1 and seen[0]["seed"] == 42
    asyncio.run(fw.run_final_writer_agent(prompts, "{}", "plan", "[]", transport=httpx.MockTransport(handler)))
    assert seen[1]["seed"] == fw.SEED


def test_final_writer_agent_passes_open_items_and_verdict():
    import asyncio

    from app import prompts
    from app.agents import final_writer as fw

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
            {"sections": [{"id": sid, "title": fw.SECTION_TITLES[sid], "text": "body."} for sid in fw.SECTION_IDS],
             "open_items": ["What is the CEO's relevant experience?", "  ", "What price per unit?"],
             "verdict": "Weak as it stands, because the founding team's track record is not evidenced."})}}]})

    out = asyncio.run(fw.run_final_writer_agent(prompts, "{}", "plan", "[]", transport=httpx.MockTransport(handler)))
    assert out["open_items"] == ["What is the CEO's relevant experience?", "What price per unit?"]
    assert out["verdict"].startswith("Weak as it stands")


def test_pipeline_rejects_non_integer_seed():
    from app import pipeline as pl
    from app.contract import ModuleError

    with pytest.raises(ModuleError):
        pl.run(Request(request_id="x", options={"answers": {"sector": "a"}, "seed": "abc"}), None)


def test_pipeline_seed_reaches_final_writer(monkeypatch):
    from app import pipeline as pl
    from app.agents import final_writer as fw
    from app.llm import LLMClient

    monkeypatch.setenv("BP_LLM_STEPS", "critic,final_writer")
    monkeypatch.setenv("LLM_API_KEY", "k")
    seen = []

    def sync_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(
            {"weaknesses": []})}}], "usage": {"prompt_tokens": 3, "completion_tokens": 2}})

    async def fake_post(client, body, remaining):
        seen.append(body)
        return json.dumps({"sections": [{"id": sid, "title": fw.SECTION_TITLES[sid], "text": "body."}
                                        for sid in fw.SECTION_IDS]}), {"prompt_tokens": 5, "completion_tokens": 4}

    monkeypatch.setattr(fw, "_post", fake_post)
    llm = LLMClient(api_key="k", transport=httpx.MockTransport(sync_handler), cache=False, max_retries=0)
    r = pl.run(Request(request_id="fp", lang="en", options={"answers": ANSWERS, "generate_plan": True, "seed": 99}), llm)
    assert r.status == "ok"
    assert len(seen) == 1 and seen[0]["seed"] == 99
