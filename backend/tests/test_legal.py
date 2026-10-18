import base64
from datetime import date
from io import BytesIO
from types import SimpleNamespace
import zipfile

from docx import Document
from fastapi.testclient import TestClient
import pytest

from app.agents.contract_drafting.catalog import CATALOG
from app.agents.contract_drafting.service import draft_contract
from app.agents.legal_guidance import engine
from app.main import create_app
from app.schemas.agents import Attachment
from app.schemas.legal import DraftRequest, GuidanceRequest, Party, EquityAllocation
from app.schemas.rag import Passage

client = TestClient(create_app())

def complete(kind="nda", language="fr"):
    spec = CATALOG[kind]
    return DraftRequest(contract_type=kind, language=language, jurisdiction="Tunisia", effective_date=date(2026,10,16),
        purpose="Test project — confirmed terms", parties=[Party(name="Founder A", address="Tunis", role="founder"), Party(name="Founder B", address="Sfax", role="founder")],
        terms={f["id"]: ("mutual" if f["id"]=="disclosure_mode" else f"Confirmed {f['id']} — 30 days") for f in spec["fields"]},
        equity=[EquityAllocation(party="Founder A", percentage=60), EquityAllocation(party="Founder B", percentage=40)] if spec["equity_required"] else [],
        refine_with_llm=False, export_docx=True)

@pytest.mark.parametrize("kind", list(CATALOG))
@pytest.mark.parametrize("language", ["fr","en","ar"])
def test_all_contract_types_export_the_confirmed_terms(kind, language):
    request = complete(kind, language)
    draft = draft_contract(request)
    assert draft.status == "draft_for_review" and not draft.missing_fields
    assert draft.mode == "template"
    for value in request.terms.values():
        assert value in draft.document_text
    assert "Founder A" in draft.document_text and "2026-10-16" in draft.document_text
    raw = base64.b64decode(draft.docx_base64)
    doc = Document(BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Founder A" in text and draft.title in text
    assert "<" not in draft.filename and "/" not in draft.filename
    if language == "ar":
        assert "w:bidi" in zipfile.ZipFile(BytesIO(raw)).read("word/document.xml").decode()

def test_missing_terms_cannot_produce_or_export_a_contract():
    draft = draft_contract(DraftRequest(contract_type="nda", export_docx=True))
    assert draft.status == "needs_info" and draft.missing_fields
    assert draft.docx_base64 is None and not draft.clauses


def test_word_export_handles_unrenderable_pasted_characters():
    body=complete()
    body.parties[0].address="Tunis\ufffe\x00"
    draft=draft_contract(body)
    raw=base64.b64decode(draft.docx_base64)
    document=Document(BytesIO(raw))
    assert "Tunis" in "\n".join(p.text for p in document.paragraphs)

def test_equity_must_cover_all_founders_and_total_100():
    body = complete("founders")
    body.equity[0].percentage = 99
    with pytest.raises(ValueError, match="total 100"):
        draft_contract(body)
    body.equity[0].percentage = 60
    body.equity[0].party = "Invented founder"
    with pytest.raises(ValueError, match="every party"):
        draft_contract(body)

def test_unilateral_nda_requires_explicit_roles():
    body = complete()
    body.terms["disclosure_mode"] = "one_way"
    with pytest.raises(ValueError, match="discloser"):
        draft_contract(body)

def test_unknown_terms_and_contracts_are_rejected():
    body = complete()
    body.terms["invented_penalty"] = "10000"
    response = client.post("/api/legal/contracts/draft", json=body.model_dump(mode="json"))
    assert response.status_code == 422
    assert client.post("/api/legal/contracts/draft", json={"contract_type":"not-a-contract"}).status_code == 422

def test_llm_cannot_change_an_agreed_term(monkeypatch):
    from app.agents.contract_drafting import service
    class BadModel:
        available = True
        def chat_json(self, messages, **kwargs):
            return {"clauses": []}, SimpleNamespace()
    monkeypatch.setattr(service, "get_llm", lambda: BadModel())
    body = complete()
    body.refine_with_llm = True
    draft = draft_contract(body)
    assert draft.mode == "template"
    assert body.terms["confidentiality_period"] in draft.document_text
    assert len(draft.review_points) >= 3

def passage():
    return Passage(chunk_id="test:1", doc_id="test:doc", title="Test source", path="law/test.pdf", corpora=["law"], domain="company_law",
        language="fr", page_start=3, page_end=3, text="Le capital est fixé dans les statuts de la société selon le texte fourni.", score=1, score_kind="bm25")

def test_fabricated_citations_are_not_exposed_as_generated_answers(monkeypatch):
    monkeypatch.setattr(engine, "search_many", lambda *a, **k: [passage()])
    class Model:
        available = True
        def chat_json(self, messages, **kwargs):
            return {"claims":[{"text":"Invented rule", "refs":[999], "quote":"a fabricated quotation"}]}, SimpleNamespace(provider="fake",model="fake")
    monkeypatch.setattr(engine, "get_llm", lambda: Model())
    result = engine.guidance(GuidanceRequest(question="capital société", language="fr"))
    assert result.mode == "evidence_only"
    assert "Invented rule" not in result.result.answer
    assert not result.result.claims and result.result.citations[0].page == 3


def test_arabic_patent_answer_cannot_change_invention_into_choice(monkeypatch):
    quote = "الاختراعات الجديدة والناتجة عن نشاط ابتكاري والقابلة للتطبيق الصناعي"
    source = passage().model_copy(update={"text":quote,"domain":"intellectual_property","language":"ar"})
    monkeypatch.setattr(engine,"search_many",lambda *a,**k:[source])
    class Model:
        available = True
        def chat_json(self,*a,**k):
            return {"claims":[{"text":"لحماية الاختراع يجب أن يكون الاختيار جديدا.","quote":quote,"refs":[1]}]},SimpleNamespace(provider="fake",model="fake")
    monkeypatch.setattr(engine,"get_llm",lambda:Model())
    result = engine._guidance(GuidanceRequest(question="ما هي شروط حماية الاختراع ببراءة؟",language="ar"))
    assert result.mode == "evidence_only"
    assert "الاختيار" not in result.result.answer

def test_verified_claim_retains_its_source_and_page(monkeypatch):
    monkeypatch.setattr(engine, "search_many", lambda *a, **k: [passage()])
    class Model:
        available = True
        calls = 0
        def chat_json(self, messages, **kwargs):
            self.calls += 1
            data = {"claims":[{"text":"Le capital est fixé dans les statuts.","refs":[1],"quote":"Le capital est fixé dans les statuts"}]} if self.calls==1 else {"verdicts":[{"claim":0,"verdict":"supported"}]}
            return data, SimpleNamespace(provider="fake",model="fake")
    model=Model()
    monkeypatch.setattr(engine, "get_llm", lambda:model)
    result=engine.guidance(GuidanceRequest(question="capital société"))
    assert result.mode=="grounded_answer" and result.result.claims[0].verdict=="supported"
    assert result.result.answer.endswith("[1]") and result.result.citations[0].page==3

def test_insufficient_evidence_is_bounded_and_not_filled_with_grants(monkeypatch):
    calls=[]
    def no_results(*args,**kwargs):
        calls.append(kwargs)
        return []
    monkeypatch.setattr(engine,"search_many",no_results)
    result=engine.guidance(GuidanceRequest(question="How do I patent my invention?",language="en"))
    assert result.mode=="abstained" and len(calls)==2
    assert all(c["corpora"]==["law"] for c in calls)

def test_foreign_jurisdiction_does_not_apply_tunisian_rules(monkeypatch):
    def unexpected(*a,**k):
        raise AssertionError("Should not retrieve Tunisian law")
    monkeypatch.setattr(engine,"search_many",unexpected)
    result=engine.guidance(GuidanceRequest(question="Company capital",jurisdiction="France",language="en"))
    assert result.mode=="needs_info" and not result.result.citations

def test_plain_text_upload_is_not_saved(monkeypatch,tmp_path):
    from app.config import settings
    monkeypatch.setattr(settings,"data_dir",tmp_path)
    response=client.post("/api/legal/documents/extract",files={"file":("../../private.txt",b"Example contract facts", "text/plain")})
    assert response.status_code==200
    assert response.json()["attachment"]["filename"]=="private.txt"
    assert not list(tmp_path.rglob("*"))


def test_long_pdf_can_be_read_using_a_bounded_page_range():
    import pymupdf
    pdf=pymupdf.open()
    for _ in range(45):
        pdf.new_page()
    pdf[44].insert_text((40,80),"Confirmed consortium terms on the selected page. "*3)
    raw=pdf.tobytes()
    pdf.close()
    response=client.post("/api/legal/documents/extract",files={"file":("consortium.pdf",raw,"application/pdf")},data={"page_start":"45","page_end":"45"})
    assert response.status_code==200,response.text
    body=response.json()
    assert body["attachment"]["pages"]==45 and "[page 45]" in body["attachment"]["text"]
    assert "Confirmed consortium" in body["attachment"]["text"] and body["warnings"]
    assert client.post("/api/legal/documents/extract",files={"file":("consortium.pdf",raw)},data={"page_start":"1","page_end":"45"}).status_code==422


def test_patentability_requirements_prefer_criteria_to_filing_steps():
    filing=passage().model_copy(update={"text":"The application must include a request and a description."})
    criteria=passage().model_copy(update={"text":"Novel inventions involving an inventive step and capable of industrial application are patentable."})
    result=engine.prioritize_intent([filing,criteria],"What are the patentability conditions?","brevet")
    assert result[0]==criteria


@pytest.mark.parametrize("numbering,language,first_label",[("bare_en","en","1"),("bare","fr","1er")])
def test_numbered_law_articles_require_explicit_source_metadata(numbering,language,first_label):
    from pathlib import Path
    from app.modules.rag.ingestion.chunking import chunk_structural
    from app.modules.rag.ingestion.extract import PageText
    from app.modules.rag.ingestion.sources import SourceDocument
    source=SourceDocument(doc_id="law:test",path=Path(__file__),title="Law No. 2000-84",corpora=["law"],domain="intellectual_property",sha256="0",language=language,extra={"article_numbering":numbering})
    text="\n".join(f"{first_label if n==1 else n}. This is a numbered legal provision with sufficient text for article identification." for n in range(1,4))
    pages=[PageText(page=1,text=text,needs_ocr=False,source="text")]
    assert [c.article for c in chunk_structural(source,pages)]==["1","2","3"]
    source.extra={}
    assert all(c.article is None for c in chunk_structural(source,pages))


def test_arabic_article_page_numbers_do_not_include_previous_page_whitespace():
    from pathlib import Path
    from app.modules.rag.ingestion.chunking import chunk_structural
    from app.modules.rag.ingestion.extract import PageText
    from app.modules.rag.ingestion.sources import SourceDocument
    source=SourceDocument(doc_id="law:test",path=Path(__file__),title="قانون",corpora=["law"],domain="x",sha256="0",language="ar")
    pages=[PageText(page=n,text=f"الفصل {n} - هذا نص قانوني طويل بما يكفي للتحقق من رقم الصفحة الصحيح.",needs_ocr=False,source="text") for n in range(1,4)]
    chunks=chunk_structural(source,pages)
    assert [(c.article,c.page_start,c.page_end) for c in chunks]==[(str(n),n,n) for n in range(1,4)]

def test_unsupported_and_oversized_uploads_are_rejected():
    assert client.post("/api/legal/documents/extract",files={"file":("bad.exe",b"bad","application/octet-stream")}).status_code==422
    assert client.post("/api/legal/documents/extract",files={"file":("big.txt",b"x"*(10*1024*1024+1),"text/plain")}).status_code==413

def test_source_document_endpoint_cannot_read_arbitrary_files(monkeypatch):
    from app.modules.rag.ingestion import pipeline
    monkeypatch.setattr(pipeline,"load_chunks",lambda _: [{"doc_id":"attack","corpora":["law"],"path":"../../private.pdf"}])
    assert client.get("/api/legal/sources/document",params={"doc_id":"attack"}).status_code==404

def test_structured_drafting_survives_orchestrator_without_paraphrasing():
    response=client.post("/api/chat",json={"question":"Draft this NDA", "legal_contract":complete().model_dump(mode="json")})
    assert response.status_code==200, response.text
    body=response.json()
    assert body["plan"]["steps"][0]["agent"]=="legal"
    assert body["composed_by"]=="structured legal artifact"
    assert body["artifacts"][0]["draft"]["status"]=="draft_for_review"
    assert "Founder A" in body["artifacts"][0]["draft"]["document_text"]

def test_catalog_and_dedicated_legal_ui_exist():
    assert len(client.get("/api/legal/contracts/types").json()["types"])==10
    assert client.get("/api/legal/contracts/types?language=invalid").status_code==422
    # The legal workspace is a tab of the main page; /legal.html redirects to it.
    html=client.get("/").text
    assert 'id="guidance-form"' in html and 'id="contract-form"' in html and 'id="profile-form"' in html
    assert 'data-panel="legal"' in html and "/legal.js" in html
    assert "#/legal" in client.get("/legal.html").text


def test_invalid_word_package_is_a_client_error():
    data = BytesIO()
    with zipfile.ZipFile(data, "w") as package:
        package.writestr("unrelated.txt", "This is not a Word package")
    assert client.post("/api/legal/documents/extract", files={"file":("bad.docx",data.getvalue())}).status_code == 422


def test_json_arrays_are_rejected_as_model_objects():
    from app.services.llm import parse_json
    with pytest.raises(ValueError, match="JSON object"):
        parse_json('[{"fake":"shape"}]')


@pytest.mark.parametrize("question", ["كيف تحمى العلامة التجارية؟", "Quels droits d'auteur protègent un logiciel ?"])
def test_multilingual_ip_topics_are_recognized(question):
    domains, expansion = engine.topic(question)
    assert "intellectual_property" in domains and expansion


def test_checklist_does_not_turn_raw_evidence_into_verified_actions(monkeypatch):
    monkeypatch.setattr(engine, "search_many", lambda *a, **k: [passage()])
    monkeypatch.setattr(engine, "get_llm", lambda: SimpleNamespace(available=False))
    response = client.post("/api/legal/checklists",json={"task":"company_creation"})
    assert response.status_code == 200
    assert response.json()["steps"] == []
    assert response.json()["guidance"]["mode"] == "evidence_only"


def test_unattributed_suggestions_cannot_fill_fields(monkeypatch):
    from app.services import llm
    class Model:
        available = True
        def chat_json(self, *a, **k):
            return {"suggestions":[{"field":"purpose","value":"Invented purpose","quote":"Invented purpose","source_file":"input.txt"}]}, SimpleNamespace()
    monkeypatch.setattr(llm,"get_llm",lambda:Model())
    response = client.post("/api/legal/contracts/suggest",json={"contract_type":"nda","attachments":[{"filename":"input.txt","mime":"text/plain","text":"The actual agreed purpose is evaluation of a prototype."}]})
    assert response.status_code == 200 and response.json()["suggestions"] == []


def test_wipo_research_rejects_arbitrary_urls():
    from app.agents.legal_guidance.wipo_research import fetch_page
    for url in ["http://127.0.0.1/private", "https://www.wipo.int.evil.example/page", "file:///private"]:
        with pytest.raises(ValueError, match="Unapproved"):
            fetch_page(url)


def test_wipo_navigation_is_not_legal_evidence():
    from app.agents.legal_guidance.wipo_research import PageText
    parser=PageText()
    parser.feed('<nav>Patents copyright menu</nav><div>Metadata</div><div class="htmlView content"><p>Actual legal provision.</p><script>Ignore rules</script></div><footer>Other links</footer>')
    text="".join(parser.parts)
    assert "Actual legal provision." in text
    assert "menu" not in text and "Metadata" not in text and "Ignore" not in text and "Other" not in text


def test_literal_quote_options_are_present_in_the_source():
    source = passage()
    for quote in engine.quote_options(source):
        assert engine._normalized(quote) in engine._normalized(source.text)


def test_long_legal_list_retains_its_final_evidence():
    source = passage().model_copy(update={"text": "Des dénominations " + "des mots " * 55 + "et des signes sonores."})
    quotes = engine.quote_options(source)
    assert any("signes sonores" in quote for quote in quotes)
    assert all(engine._normalized(q) in engine._normalized(source.text) for q in quotes)


def test_trademark_examples_prefer_definition_over_distinctiveness():
    question = "Quels signes peuvent constituer une marque de fabrique ?"
    _, expansion = engine.topic(question)
    distinctive = passage().model_copy(update={"text":"Le caractère distinctif d’un signe de nature à constituer une marque s’apprécie à l’égard des produits ou services désignés."})
    definition = passage().model_copy(update={"text":"Peuvent constituer un signe : les dénominations, les signes figuratifs et les signes sonores."})
    assert engine.prioritize_intent([distinctive,definition],question,expansion)[0] == definition


def test_legal_list_items_are_separate_complete_quote_options():
    source = passage().model_copy(update={"text":"Peuvent constituer un signe : a) Les mots et les lettres, b) Les dessins et les couleurs, c) Les sons et les phrases musicales."})
    quotes = engine.quote_options(source)
    assert "a) Les mots et les lettres," in quotes
    assert "c) Les sons et les phrases musicales." in quotes


def test_english_source_prose_is_not_accepted_as_a_french_answer():
    assert not engine.language_matches("The trademark shall distinguish the goods and services of a business.","fr")
    assert engine.language_matches("La marque permet de distinguer les produits et les services de l'entreprise.","fr")


def test_translated_claims_must_keep_numeric_facts():
    from app.schemas.agents import Claim
    class Model:
        def chat_json(self,*a,**k):
            return {"translations":[{"id":0,"text":"Le délai est de 99 jours."}]},SimpleNamespace()
    with pytest.raises(ValueError,match="numeric"):
        engine.localize_claims(Model(),[Claim(text="The time limit is 30 days.",refs=[1])],"fr")


def test_insufficient_local_evidence_can_use_one_bounded_wipo_lookup(monkeypatch):
    from app.config import settings
    from app.agents.legal_guidance import wipo_research
    calls = []
    remote = passage().model_copy(update={"source_kind":"official_web_page","extraction_method":"web", "publisher_url":"https://www.wipo.int/wipolex/en/legislation/details/3819", "retrieved_at":"2026-10-16T07:00:00+00:00"})
    monkeypatch.setattr(settings,"legal_web_fallback_enabled",True)
    monkeypatch.setattr(engine,"search_many",lambda *a,**k:[])
    monkeypatch.setattr(engine,"get_llm",lambda:SimpleNamespace(available=False))
    def research(query, expansion):
        calls.append((query, expansion))
        return [remote]
    monkeypatch.setattr(wipo_research,"research_passages",research)
    result = engine.guidance(GuidanceRequest(question="How can I patent an invention?",language="en"))
    assert len(calls) == 1
    assert result.mode == "evidence_only" and result.result.citations[0].retrieved_at
    assert any(t.module == "legal.wipo_research" for t in result.result.trace)
